import openpyxl,json,pathlib,shutil,sys
source=pathlib.Path(sys.argv[1])
root=pathlib.Path(__file__).resolve().parents[1]
(root/'public').mkdir(exist_ok=True)
shutil.copyfile(source,root/'public/template.xlsx')
w=openpyxl.load_workbook(source)
# Deliberate entry ranges, checked against every source sheet. Never infer inputs from all blank cells.
ranges=[['B15:AC21'],['B6:V36'],['C5:K24'],['B7:R24'],['C7:N8','C13:N14','C18:N19','C24:N25','C29:N30'],['B3:K22'],['B7:H21','C3:C4','G3:G4'],['B7:G21']]
bounds=[(30,29),(41,22),(27,11),(28,18),(33,15),(23,11),(30,8),(30,7)]
headercells=[['A2'],['B1'],['A1'],['B1'],['B2','B4','B15','B21'],[],[],[]]
def color(c,default):
    return '#'+c.rgb[-6:] if c and c.type=='rgb' else default
sheets=[]
for idx,s in enumerate(w):
    rows,cols=bounds[idx]
    merged={}
    skip=set()
    for m in s.merged_cells.ranges:
        merged[(m.min_row,m.min_col)]=(m.max_row-m.min_row+1,m.max_col-m.min_col+1)
        skip.update((r,c) for r in range(m.min_row,m.max_row+1) for c in range(m.min_col,m.max_col+1) if (r,c)!=(m.min_row,m.min_col))
    editable=set()
    for area in ranges[idx]:
        for row in s[area]:
            for c in row:
                if not isinstance(c,openpyxl.cell.cell.MergedCell): editable.add(c.coordinate)
    widths=[]
    for c in range(1,cols+1):
        width=s.sheet_format.defaultColWidth or 8.43
        for d in s.column_dimensions.values():
            if d.min<=c<=d.max: width=d.width
        widths.append(round(width*7+5,2))
    cells=[]
    for row in s.iter_rows(min_row=1,max_row=rows,max_col=cols):
        for c in row:
            if (c.row,c.column) in skip: continue
            rs,cs=merged.get((c.row,c.column),(1,1))
            borders={}
            for side in ['top','bottom','left','right']:
                b=getattr(c.border,side)
                if b and b.style: borders[side]=('2px' if b.style in ['medium','thick','double'] else '1px')+' '+('dashed' if 'dash' in b.style.lower() else 'solid')+' '+color(b.color,'#111111')
            cells.append(dict(ref=c.coordinate,row=c.row,col=c.column,rs=rs,cs=cs,text='' if c.value is None else str(c.value),editable=c.coordinate in editable,header=c.coordinate in headercells[idx],font=c.font.name,size=c.font.sz or 11,bold=bool(c.font.b),align=c.alignment.horizontal or 'center',valign=c.alignment.vertical or 'center',wrap=bool(c.alignment.wrap_text),color=color(c.font.color,'#111111'),fill=color(c.fill.fgColor,'#ffffff') if c.fill.patternType else '#ffffff',borders=borders))
    sheets.append(dict(id=str(idx+1),name=s.title,rows=rows,cols=cols,widths=widths,heights=[s.row_dimensions[r].height or s.sheet_format.defaultRowHeight or 15 for r in range(1,rows+1)],cells=cells,orientation=s.page_setup.orientation or 'portrait',scale=s.page_setup.scale or 100,printArea=str(s.print_area),margins={k:getattr(s.page_margins,k) for k in ['left','right','top','bottom']}))
(root/'src').mkdir(exist_ok=True)
(root/'src/workbook.json').write_text(json.dumps(dict(version=1,sheets=sheets),ensure_ascii=False),encoding='utf-8')
print(json.dumps([dict(name=s['name'],fields=sum(c['editable'] for c in s['cells'])) for s in sheets],ensure_ascii=False))
